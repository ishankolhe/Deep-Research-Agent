import os
import re
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE

from utils.markdown_parse import parse_markdown

ACCENT = RGBColor(0x2E, 0x31, 0x92)
ACCENT_LIGHT = RGBColor(0x6C, 0x6F, 0xC4)
INK = RGBColor(0x1A, 0x1A, 0x2E)
MUTED = RGBColor(0x6B, 0x72, 0x80)
REF_RE = re.compile(r"^\[(\d+)\]\s+(.*)$")


def _add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _add_hyperlink(paragraph, url: str, text: str, color: str = "2E3192"):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    rpr.append(color_el)
    u_el = OxmlElement("w:u")
    u_el.set(qn("w:val"), "single")
    rpr.append(u_el)
    sz_el = OxmlElement("w:sz")
    sz_el.set(qn("w:val"), "18")
    rpr.append(sz_el)
    new_run.append(rpr)
    t_el = OxmlElement("w:t")
    t_el.text = text
    new_run.append(t_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _shade_cell(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _add_runs(paragraph, runs):
    for text, bold, citation in runs:
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = bold
        if citation:
            run.font.superscript = True
            run.font.size = Pt(9)
            run.font.color.rgb = ACCENT


def _add_bottom_border(paragraph, color="2E3192", size=18):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def export_docx(query: str, markdown: str, chart_paths: list[str], image_paths: list[str],
                 output_path: str) -> str:
    doc = Document()

    # Base typography
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK

    section = doc.sections[0]
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # ---- Title page ----
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kicker.add_run("DEEP RESEARCH REPORT")
    kr.font.size = Pt(12)
    kr.font.color.rgb = ACCENT_LIGHT
    kr.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(query.strip().capitalize())
    tr.font.size = Pt(28)
    tr.bold = True
    tr.font.color.rgb = INK

    rule = doc.add_paragraph()
    _add_bottom_border(rule)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(f"Generated {datetime.now().strftime('%B %d, %Y')} · AI Deep Research Agent")
    mr.font.size = Pt(10)
    mr.font.color.rgb = MUTED
    mr.italic = True

    doc.add_page_break()

    # ---- Footer with page number ----
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Deep Research Agent  |  Page ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED
    _add_page_number_field(footer_p)

    # ---- Body ----
    blocks = parse_markdown(markdown)
    ref_mode = False
    for block in blocks:
        btype = block["type"]

        if btype == "h1":
            p = doc.add_paragraph()
            r = p.add_run(block["text"])
            r.font.size = Pt(20)
            r.bold = True
            r.font.color.rgb = ACCENT
            p.space_before = Pt(6)

        elif btype == "h2":
            ref_mode = block["text"].strip().lower() == "references"
            p = doc.add_paragraph()
            p.space_before = Pt(18)
            p.space_after = Pt(6)
            r = p.add_run(block["text"])
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = ACCENT
            _add_bottom_border(p, color="D8D8E8", size=6)

        elif btype == "h3":
            p = doc.add_paragraph()
            p.space_before = Pt(12)
            r = p.add_run(block["text"])
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = ACCENT_LIGHT

        elif btype == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, block["runs"])

        elif btype == "numbered":
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, block["runs"])

        elif btype == "table":
            rows = block["rows"]
            if not rows:
                continue
            ncols = len(rows[0])
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Light Grid Accent 1"
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row[:ncols]):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    clean = re.sub(r"\*\*(.*?)\*\*", r"\1", cell_text)
                    run = para.add_run(clean)
                    run.font.size = Pt(10)
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        _shade_cell(cell, "2E3192")
            doc.add_paragraph()

        else:  # plain paragraph
            if ref_mode:
                text = "".join(t for t, _, _ in block["runs"])
                m = REF_RE.match(text.strip())
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                if m:
                    num, url = m.group(1), m.group(2)
                    lead = p.add_run(f"[{num}]  ")
                    lead.font.size = Pt(9)
                    lead.font.color.rgb = MUTED
                    _add_hyperlink(p, url, url)
                else:
                    r = p.add_run(text)
                    r.font.size = Pt(9)
                    r.font.color.rgb = MUTED
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                _add_runs(p, block["runs"])

    # ---- Visuals section ----
    if chart_paths or image_paths:
        doc.add_page_break()
        h = doc.add_paragraph()
        h.space_after = Pt(10)
        r = h.add_run("Charts & Visual References")
        r.font.size = Pt(18)
        r.bold = True
        r.font.color.rgb = ACCENT
        _add_bottom_border(h)

        for i, p_path in enumerate(chart_paths):
            if os.path.exists(p_path):
                doc.add_picture(p_path, width=Inches(5.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cap.add_run(f"Figure {i + 1} — data derived from research findings")
                cr.font.size = Pt(9)
                cr.italic = True
                cr.font.color.rgb = MUTED
                cap.paragraph_format.space_after = Pt(14)

        if image_paths:
            sub = doc.add_paragraph()
            sub.space_before = Pt(8)
            sr = sub.add_run("Illustrative Images")
            sr.font.size = Pt(13)
            sr.bold = True
            sr.font.color.rgb = ACCENT_LIGHT
            for p_path in image_paths:
                if os.path.exists(p_path):
                    doc.add_picture(p_path, width=Inches(4.6))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(output_path)
    return output_path
